"""Đọc 3 file XLSX mẫu → JSON cho test/e2e.test.js (test chạy trên ĐÚNG file sẽ upload lên Drive).

Chạy: python3 tools/export_fixture.py [thư_mục_xlsx] [file_json]
      mặc định: templates/ → test/fixtures/templates.json
Ô ngày → {"$date": "YYYY-MM-DDTHH:MM:SS"} · ô công thức → {"$f": "=..."}
"""
import openpyxl, json, os, sys, datetime

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'templates')
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, 'test', 'fixtures', 'templates.json')
FILES = {'MAIN': 'THAIPUT_MAIN_v6.1.xlsx', 'TUTOR': 'THAIPUT_TUTOR_v6.1.xlsx', 'REGISTRATION': 'THAIPUT_REGISTRATION_v6.1.xlsx'}

def cell(v):
    if v is None: return ''
    if isinstance(v, (datetime.datetime, datetime.date)): return {'$date': v.isoformat()}
    if isinstance(v, str) and v.startswith('='): return {'$f': v}
    return v

out = {}
for key, fn in FILES.items():
    wb = openpyxl.load_workbook(os.path.join(SRC, fn))
    out[key] = {}
    for ws in wb.worksheets:
        rows = [[cell(c) for c in r] for r in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column, values_only=True)]
        while rows and all(v == '' for v in rows[-1]): rows.pop()
        out[key][ws.title] = rows
with open(OUT, 'w', encoding='utf-8') as f:
    json.dump(out, f, ensure_ascii=False, indent=1, sort_keys=False)
    f.write('\n')
print('Wrote', OUT)
